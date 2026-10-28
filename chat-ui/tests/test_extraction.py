"""Tests for CV document text extraction.

These build real PDF, Word and OpenDocument files in memory (no fixtures on
disk) so the extraction code is exercised against genuine file bytes.
"""

import io
import zipfile

import pytest
from docx import Document
from docx.oxml import parse_xml
from reportlab.pdfgen import canvas

from extraction import extract_text


def _make_docx(*paragraphs: str) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_pdf(*lines: str) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    y = 800
    for line in lines:
        c.drawString(72, y, line)
        y -= 20
    c.save()
    return buf.getvalue()


def test_extracts_text_from_docx():
    data = _make_docx("Jane Doe", "Senior Engineer at Acme")
    text = extract_text("cv.docx", data)
    assert "Jane Doe" in text
    assert "Senior Engineer at Acme" in text


def test_extracts_text_from_pdf():
    data = _make_pdf("John Smith", "Product Manager")
    text = extract_text("resume.pdf", data)
    assert "John Smith" in text
    assert "Product Manager" in text


def test_extension_match_is_case_insensitive():
    data = _make_docx("Hello")
    assert "Hello" in extract_text("CV.DOCX", data)


def test_unsupported_extension_raises_valueerror():
    with pytest.raises(ValueError) as exc:
        extract_text("cv.png", b"not text")
    assert "png" in str(exc.value).lower()


def test_unreadable_pdf_raises_valueerror():
    with pytest.raises(ValueError):
        extract_text("broken.pdf", b"not really a pdf")


# --- Word ------------------------------------------------------------------

def _docx_bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_includes_table_text_in_document_order():
    doc = Document()
    doc.add_paragraph("Experience")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Engineer, Acme"
    table.cell(0, 1).text = "2019 - 2021"
    doc.add_paragraph("Education")
    text = extract_text("cv.docx", _docx_bytes(doc))
    assert text.index("Experience") < text.index("Engineer, Acme") < text.index("Education")


def test_docx_joins_single_line_cells_of_a_row():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Engineer, Acme"
    table.cell(0, 1).text = "2019 - 2021"
    assert "Engineer, Acme | 2019 - 2021" in extract_text("cv.docx", _docx_bytes(doc))


def test_docx_layout_table_keeps_each_column_together():
    # A two-column template: a sidebar cell and a main cell, several lines each.
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    side, main = table.cell(0, 0), table.cell(0, 1)
    side.text = "Skills"
    side.add_paragraph("Python")
    main.text = "Experience"
    main.add_paragraph("Engineer at Acme")
    lines = extract_text("cv.docx", _docx_bytes(doc)).splitlines()
    assert lines == ["Skills", "Python", "Experience", "Engineer at Acme"]


def test_docx_merged_cells_are_not_repeated():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Jane Doe"
    assert extract_text("cv.docx", _docx_bytes(doc)).count("Jane Doe") == 1


def test_docx_includes_header_and_footer_text():
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "jane@example.com"
    doc.add_paragraph("Profile")
    doc.sections[0].footer.paragraphs[0].text = "References on request"
    text = extract_text("cv.docx", _docx_bytes(doc))
    assert text.index("jane@example.com") < text.index("Profile") < text.index("References on request")


# A floating text box as Word writes it: the shape (mc:Choice) plus a VML copy
# of the same text (mc:Fallback) for older readers.
_TEXT_BOX_RUN = (
    '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    ' xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
    ' xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
    ' xmlns:v="urn:schemas-microsoft-com:vml">'
    "<mc:AlternateContent>"
    '<mc:Choice Requires="wps"><w:drawing><wps:wsp><wps:txbx><w:txbxContent>'
    "<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"
    "</w:txbxContent></wps:txbx></wps:wsp></w:drawing></mc:Choice>"
    "<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent>"
    "<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"
    "</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback>"
    "</mc:AlternateContent></w:r>"
)


def test_docx_includes_text_box_text_once():
    doc = Document()
    p = doc.add_paragraph("Profile")
    p._p.append(parse_xml(_TEXT_BOX_RUN.format(text="07700 900123")))
    text = extract_text("cv.docx", _docx_bytes(doc))
    assert text.count("07700 900123") == 1
    assert "Profile" in text


def test_docx_keeps_tabs_and_line_breaks():
    doc = Document()
    p = doc.add_paragraph("Engineer")
    p.add_run().add_tab()
    p.add_run("2019")
    p.add_run().add_break()
    p.add_run("Acme")
    assert "Engineer\t2019\nAcme" in extract_text("cv.docx", _docx_bytes(doc))


# --- PDF -------------------------------------------------------------------

def _make_pdf_at(*items: tuple[float, float, str], links=()) -> bytes:
    """Draw each (x, y, text) on one A4 page; links are (url, rect) pairs."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for x, y, line in items:
        c.drawString(x, y, line)
    for url, rect in links:
        c.linkURL(url, rect)
    c.save()
    return buf.getvalue()


def test_pdf_two_columns_are_read_one_after_the_other():
    left = ["Skills", "Python", "SQL", "Excel"]
    right = ["Experience", "Engineer at Acme", "Analyst at Globex", "Intern at Initech"]
    # Drawn row by row, as most generators do, so the columns interleave.
    items = [(x, 700 - i * 20, t) for i, pair in enumerate(zip(left, right)) for x, t in zip((50, 300), pair)]
    lines = extract_text("cv.pdf", _make_pdf_at(*items)).splitlines()
    assert lines == left + right


def test_pdf_full_width_header_and_footer_stay_outside_the_columns():
    left = ["Skills", "Python", "SQL", "Excel"]
    right = ["Experience", "Engineer at Acme", "Analyst at Globex", "Intern at Initech"]
    items = [(200, 780, "Jane Doe - Data Analyst - London")]
    items += [(x, 700 - i * 20, t) for i, pair in enumerate(zip(left, right)) for x, t in zip((50, 300), pair)]
    items += [(250, 50, "Page 1 of 1")]
    lines = extract_text("cv.pdf", _make_pdf_at(*items)).splitlines()
    assert lines == ["Jane Doe - Data Analyst - London", *left, *right, "Page 1 of 1"]


def test_pdf_single_column_with_right_aligned_dates_stays_line_by_line():
    bullet = "Led a team of five analysts building weekly reporting for the sales team"
    items = [
        (72, 700, "Analyst, Globex"), (450, 700, "2021 - 2024"),
        (72, 680, bullet),
        (72, 660, "Engineer, Acme"), (450, 660, "2019 - 2021"),
        (72, 640, bullet),
    ]
    lines = extract_text("cv.pdf", _make_pdf_at(*items)).splitlines()
    assert lines[0] == "Analyst, Globex 2021 - 2024"
    assert lines[2] == "Engineer, Acme 2019 - 2021"


def test_pdf_appends_link_targets_missing_from_the_text():
    data = _make_pdf_at(
        (72, 700, "LinkedIn"), (72, 680, "https://example.com/portfolio"),
        links=[
            ("https://www.linkedin.com/in/janedoe", (72, 695, 140, 712)),
            ("https://example.com/portfolio", (72, 675, 260, 692)),
            ("mailto:jane@example.com", (72, 655, 200, 672)),
        ],
    )
    text = extract_text("cv.pdf", data)
    assert "https://www.linkedin.com/in/janedoe" in text
    assert "jane@example.com" in text
    assert "mailto:" not in text
    assert text.count("https://example.com/portfolio") == 1


# --- Plain text and RTF ----------------------------------------------------

@pytest.mark.parametrize("filename", ["cv.txt", "cv.md"])
def test_extracts_utf8_text(filename):
    assert extract_text(filename, "Zoë Smith\nEngineer".encode()) == "Zoë Smith\nEngineer"


def test_text_byte_order_mark_is_dropped():
    assert extract_text("cv.txt", "﻿Jane".encode("utf-8")) == "Jane"


def test_text_falls_back_to_windows_encoding():
    assert extract_text("cv.txt", "Café manager".encode("cp1252")) == "Café manager"


def test_extracts_text_from_rtf():
    rtf = rb"{\rtf1\ansi{\fonttbl\f0\fswiss Helvetica;}\f0\pard Jane Doe\par Engineer at Acme\par}"
    text = extract_text("cv.rtf", rtf)
    assert "Jane Doe" in text
    assert "Engineer at Acme" in text


# --- OpenDocument ----------------------------------------------------------

_ODF_NS = (
    'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
    'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"'
)


def _make_odt(body: str, master: str = "") -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        z.writestr(
            "content.xml",
            f"<office:document-content {_ODF_NS}><office:body><office:text>"
            f"{body}</office:text></office:body></office:document-content>",
        )
        z.writestr(
            "styles.xml",
            f"<office:document-styles {_ODF_NS}><office:master-styles>"
            f"<style:master-page>{master}</style:master-page>"
            f"</office:master-styles></office:document-styles>",
        )
    return buf.getvalue()


def test_extracts_odt_headings_paragraphs_and_lists():
    data = _make_odt(
        "<text:h>Jane Doe</text:h>"
        "<text:p>Engineer at <text:span>Acme</text:span></text:p>"
        "<text:list><text:list-item><text:p>Python</text:p></text:list-item></text:list>"
    )
    assert extract_text("cv.odt", data).splitlines() == ["Jane Doe", "Engineer at Acme", "Python"]


def test_odt_expands_spaces_tabs_and_line_breaks():
    data = _make_odt(
        '<text:p>A<text:s text:c="3"/>B<text:tab/>C<text:line-break/>D</text:p>'
    )
    assert extract_text("cv.odt", data) == "A   B\tC\nD"


def test_odt_includes_tables_and_page_header():
    data = _make_odt(
        "<table:table><table:table-row>"
        "<table:table-cell><text:p>Engineer, Acme</text:p></table:table-cell>"
        "<table:table-cell><text:p>2019 - 2021</text:p></table:table-cell>"
        "</table:table-row></table:table>",
        master="<style:header><text:p>jane@example.com</text:p></style:header>",
    )
    assert extract_text("cv.odt", data).splitlines() == [
        "jane@example.com",
        "Engineer, Acme | 2019 - 2021",
    ]


def test_unreadable_odt_raises_valueerror():
    with pytest.raises(ValueError):
        extract_text("cv.odt", b"not a zip")
